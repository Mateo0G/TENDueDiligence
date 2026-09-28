"""One-off deep XML inspection of the reference docx's exact formatting:
table borders/shading, run fonts/colors, and the cover-page/disclaimer
paragraph properties -- needed now that Phase 4 must replicate this exactly
per the Phase 0 "match exactly" styling decision, not just structurally.
"""
import docx
from docx.oxml.ns import qn

PATH = "reference/AccuBreath_Due_Diligence_Report.docx"
d = docx.Document(PATH)


def dump_run_props(run):
    rPr = run._element.find(qn("w:rPr"))
    if rPr is None:
        return "{}"
    parts = []
    for child in rPr:
        tag = child.tag.split("}")[-1]
        attrs = {k.split("}")[-1]: v for k, v in child.attrib.items()}
        parts.append(f"{tag}={attrs}")
    return " ".join(parts)


print("--- Run-level formatting samples ---")
for label, para in [
    ("Title (para 15)", d.paragraphs[15]),
    ("Subtitle (para 16)", d.paragraphs[16]),
    ("Heading1 (para 38 'Executive Summary')", d.paragraphs[38]),
    ("Normal body (para 42)", d.paragraphs[42]),
    ("Footer", d.sections[0].footer.paragraphs[0]),
]:
    print(f"\n{label}: {len(para.runs)} runs")
    for i, r in enumerate(para.runs[:3]):
        print(f"  run{i} text={r.text[:40]!r} font.name={r.font.name} font.size={r.font.size} "
              f"bold={r.font.bold} color={r.font.color.rgb if r.font.color and r.font.color.type else None}")
        print(f"    rPr xml: {dump_run_props(r)}")

print("\n--- Table 1 (Market Size Analysis) full XML formatting ---")
t = d.tables[0]
tbl = t._tbl
tblPr = tbl.find(qn("w:tblPr"))
print("tblPr xml:", tblPr.xml[:2000] if tblPr is not None else None)

print("\n--- Table 1 first row cell formatting (borders/shading) ---")
first_row = t.rows[0]
for i, cell in enumerate(first_row.cells[:2]):
    tcPr = cell._tc.find(qn("w:tcPr"))
    print(f"cell {i} tcPr xml:", tcPr.xml[:1500] if tcPr is not None else None)
    for p in cell.paragraphs:
        for r in p.runs:
            print(f"  run text={r.text!r} bold={r.font.bold} rPr={dump_run_props(r)}")

print("\n--- Column widths (Table 1) ---")
for i, col in enumerate(t.columns):
    print(f"col {i} width: {col.width}")

print("\n--- Section/page default font (styles.xml docDefaults) ---")
styles_element = d.styles.element
docDefaults = styles_element.find(qn("w:docDefaults"))
if docDefaults is not None:
    print(docDefaults.xml[:1500])
