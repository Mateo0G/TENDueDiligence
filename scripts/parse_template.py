"""Parse reference/AccuBreath_Due_Diligence_Report.docx into structured JSON.

Mechanical extraction only: heading outline, style usage, table positions
(anchored to the nearest preceding Heading 1), page setup, and header/footer
text. The editorial write-up (front-matter walkthrough, flags, open
questions) lives in docs/template_structure.md and is built from this
output plus a manual reading of the document.
"""
import json

import docx
from docx.oxml.ns import qn

SRC = "reference/AccuBreath_Due_Diligence_Report.docx"
OUT = "docs/template_structure.json"


def style_of(p_elem):
    pPr = p_elem.find(qn("w:pPr"))
    if pPr is None:
        return None
    pStyle = pPr.find(qn("w:pStyle"))
    return pStyle.get(qn("w:val")) if pStyle is not None else None


def text_of(elem):
    return "".join(n.text or "" for n in elem.iter(qn("w:t")))


def main():
    d = docx.Document(SRC)
    body = d.element.body

    headings = []
    tables = []
    current_h1 = None
    for child in body.iterchildren():
        tag = child.tag.split("}")[-1]
        if tag == "p":
            style = style_of(child)
            if style == "Heading1":
                text = text_of(child).strip()
                headings.append({"text": text, "is_spacer": text == ""})
                if text:
                    current_h1 = text
        elif tag == "tbl":
            trs = child.findall(qn("w:tr"))
            header_row = []
            if trs:
                for tc in trs[0].findall(qn("w:tc")):
                    header_row.append(text_of(tc).strip())
            tables.append(
                {
                    "section": current_h1,
                    "rows": len(trs),
                    "cols": len(header_row),
                    "header_row": header_row,
                }
            )

    style_counts = {}
    for p in d.paragraphs:
        style_counts[p.style.name] = style_counts.get(p.style.name, 0) + 1

    sec = d.sections[0]
    result = {
        "paragraph_count": len(d.paragraphs),
        "table_count": len(d.tables),
        "style_counts": style_counts,
        "headings": headings,
        "tables": tables,
        "inline_shape_count": len(d.inline_shapes),
        "page_setup": {
            "page_width_emu": sec.page_width,
            "page_height_emu": sec.page_height,
            "left_margin_emu": sec.left_margin,
            "right_margin_emu": sec.right_margin,
            "top_margin_emu": sec.top_margin,
            "bottom_margin_emu": sec.bottom_margin,
        },
        "footer_paragraphs": [p.text for p in sec.footer.paragraphs],
        "header_paragraphs": [p.text for p in sec.header.paragraphs],
        "has_toc_field": any(
            "TOC \\h" in (n.text or "") for n in d.element.iter(qn("w:instrText"))
        ),
    }

    with open(OUT, "w", encoding="utf-8") as f:
        json.dump(result, f, indent=2, ensure_ascii=False)

    print(f"Wrote extraction to {OUT}: {len(headings)} headings, {len(tables)} tables")


if __name__ == "__main__":
    main()
