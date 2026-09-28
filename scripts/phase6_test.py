"""Phase 6 proof: run Gap Analysis -> Mitigation Strategy (Stage 5) and the
final render (Stage 6) against the Solara Dx job that already has Stage 4
QA scores from Phase 5, plus the two stage-gating failure paths.

Needs real Postgres, R2, and ANTHROPIC_API_KEY -- run via:
  railway run --service worker -- <venv-python> scripts/phase6_test.py
"""
import io

import docx
from fastapi.testclient import TestClient

from app import repo
from app.main import app
from app.worker import run_one_task

client = TestClient(app)

JOB_ID = "70ab7c4f-015e-4f5f-be49-2f5df083731b"  # Solara Dx, scored in phase5_test.py
OUT_PATH = (
    r"C:\Users\mateo\AppData\Local\Temp\claude\c--Users-mateo-TENCapitalProjects-TEN-Due-Diligence"
    r"\2815135e-3d1d-4e86-bce1-695aff7ff4d8\scratchpad\phase6_final_report.docx"
)


def drain():
    n = 0
    while run_one_task():
        n += 1
    return n


def main():
    job = repo.get_job(JOB_ID)
    assert job is not None
    report_card = next(
        (s for s in repo.list_qa_scores(JOB_ID) if s["assessment_name"] == "Report Card of scores"),
        None,
    )
    assert report_card is not None, "expected Solara Dx to already have a Report Card from Phase 5"
    print(f"using job {JOB_ID} ({job['company_name']}); Report Card average={report_card['findings']['overall_average']}")

    # Gating failure paths, on fresh jobs with no QA scores / no gap analysis.
    no_qa_job = repo.create_job("Phase6 No-QA Test Co")
    resp = client.post(f"/jobs/{no_qa_job}/stage5")
    assert resp.status_code == 409, resp.text
    print("PASS -- stage5 refuses a job with no Report Card (409)")

    no_gap_job = repo.create_job("Phase6 No-Gap Test Co")
    resp = client.post(f"/jobs/{no_gap_job}/stage6")
    assert resp.status_code == 409, resp.text
    print("PASS -- stage6 refuses a job with no mitigation strategy (409)")

    # Idempotency for repeat runs of this script.
    from app.db import pool as _pool

    with _pool.connection() as conn:
        conn.execute("DELETE FROM tasks WHERE job_id=%s AND stage IN (5, 6)", (JOB_ID,))
        conn.execute(
            "UPDATE jobs SET gap_analysis=NULL, mitigation_strategy=NULL, status='awaiting_review' WHERE id=%s",
            (JOB_ID,),
        )
    print("cleared any prior Stage 5/6 state for this job")

    resp = client.post(f"/jobs/{JOB_ID}/stage5")
    assert resp.status_code == 200, resp.text
    stage5_ids = resp.json()["task_ids"]
    assert len(stage5_ids) == 2
    print(f"created {len(stage5_ids)} Stage 5 tasks")

    n = drain()
    print(f"drained {n} task(s)")

    tasks = repo.list_tasks_for_job(JOB_ID)
    stage5_tasks = [t for t in tasks if t["stage"] == 5]
    failed = [t for t in stage5_tasks if t["status"] == "failed"]
    assert not failed, f"Stage 5 tasks failed: {failed}"
    assert all(t["status"] == "completed" for t in stage5_tasks)
    print("PASS -- gap_analysis and mitigation_strategy tasks both completed")

    job = repo.get_job(JOB_ID)
    gap_analysis = job["gap_analysis"]
    mitigation_strategy = job["mitigation_strategy"]
    assert gap_analysis and gap_analysis["gaps"], "expected a non-empty gap analysis"
    assert mitigation_strategy and mitigation_strategy["mitigations"], "expected non-empty mitigations"
    print(f"PASS -- gap analysis found {len(gap_analysis['gaps'])} gaps, "
          f"mitigation strategy proposed {len(mitigation_strategy['mitigations'])} mitigations")

    gap_titles = {g["title"] for g in gap_analysis["gaps"]}
    mitigation_titles = {m["gap_title"] for m in mitigation_strategy["mitigations"]}
    unmatched = mitigation_titles - gap_titles
    assert not unmatched, f"mitigations reference gap titles that don't exist: {unmatched}"
    print("PASS -- every mitigation's gap_title matches an actual identified gap")

    # Stage 6: final render.
    resp = client.post(f"/jobs/{JOB_ID}/stage6")
    assert resp.status_code == 200, resp.text
    final_task_id = resp.json()["task_id"]
    n = drain()
    print(f"drained {n} task(s) for Stage 6")

    final_task = repo.get_task(final_task_id)
    assert final_task["status"] == "completed", f"final_render failed: {final_task}"

    job = repo.get_job(JOB_ID)
    assert job["status"] == "completed", f"expected job status completed, got {job['status']}"
    print(f"PASS -- job status is completed, output_key={job['output_key']}")

    resp = client.get(f"/jobs/{JOB_ID}/report.docx")
    assert resp.status_code == 200, resp.text
    with open(OUT_PATH, "wb") as f:
        f.write(resp.content)
    print(f"PASS -- downloaded final report.docx ({len(resp.content)} bytes)")

    reopened = docx.Document(io.BytesIO(resp.content))
    headings = [p.text for p in reopened.paragraphs if p.style.name == "Heading 1"]
    assert "Due Diligence Assessment Report Card" in headings
    assert "Gap Analysis & Mitigation Strategy" in headings
    assert len([h for h in headings if h.strip()]) >= 25 + 2, "expected the 25 body headings plus appendices"
    print(f"PASS -- final docx has {len(headings)} headings including both appendices")

    # doc.paragraphs excludes text inside table cells -- the gap titles are
    # rendered in a table (_add_table_block), so include table cell text too.
    full_text = "\n".join(p.text for p in reopened.paragraphs)
    for table in reopened.tables:
        for row in table.rows:
            for cell in row.cells:
                full_text += "\n" + cell.text
    assert str(gap_analysis["gaps"][0]["title"]) in full_text, "first gap's title not found in rendered docx"
    print("PASS -- gap analysis content confirmed present in rendered docx")

    print("\nAll Phase 6 checks passed.")


if __name__ == "__main__":
    try:
        main()
    finally:
        from app.db import pool

        pool.close()
